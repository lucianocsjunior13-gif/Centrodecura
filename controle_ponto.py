import sqlite3
from datetime import datetime
import os
import sys
import barcode
from barcode.writer import ImageWriter
from PIL import Image
import customtkinter as ctk
from tkinter import messagebox
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
def exportar_excel_personalizado(texto_edicao=""):
    desktop = obter_pasta_desktop()
    pasta_destino = os.path.join(desktop, "relatorios")
    if not os.path.exists(pasta_destino):
        os.makedirs(pasta_destino)

    tag_edicao = texto_edicao.strip() if texto_edicao.strip() else "edição#"
    tag_edicao_limpa = "".join(c for c in tag_edicao if c.isalnum() or c in ('#', '_', '-'))
    
    # Data com barras para o texto visual e hífens para o nome físico do arquivo
    data_formatada_barras = datetime.now().strftime('%d/%m/%Y')
    data_arquivo = datetime.now().strftime('%d-%m-%Y')
    
    # Nome seguro para o sistema operacional Windows
    nome_arquivo = f"{tag_edicao_limpa}_{data_arquivo}.xlsx"
    caminho_final = os.path.join(pasta_destino, nome_arquivo)

    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT f.nome, f.codigo_barra, r.data_hora
        FROM registros r
        JOIN funcionarios f ON r.funcionario_id = f.id
        ORDER BY r.id ASC
    ''')
    registros = cursor.fetchall()
    conn.close()

    total_cadastrados, total_hoje = obter_estatisticas()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controle de Ponto"
    ws.views.sheetView[0].showGridLines = True

    COR_VERDE_BANNER = "1F4E24"
    COR_VERDE_FAIXA  = "2E933C"
    COR_VERDE_ZEBRA  = "B8DABA"
    COR_BRANCO       = "FFFFFF"

    ws.row_dimensions[1].height = 48
    ws.row_dimensions[2].height = 6
    ws.row_dimensions[3].height = 28

    # Título do Banner no formato solicitado: edição#n_dd/mm/aaaa
    ws.merge_cells('C1:F1')
    titulo_cell = ws['C1']
    titulo_cell.value = f"Controle de Ponto — {tag_edicao_limpa}_{data_formatada_barras}"
    titulo_cell.font = Font(name="Segoe UI", size=14, bold=True, color=COR_BRANCO)
    titulo_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 7):
        ws.cell(row=1, column=col).fill = PatternFill(start_color=COR_VERDE_BANNER, end_color=COR_VERDE_BANNER, fill_type="solid")

    logo_path = obter_caminho_recurso("logocura2.png")
    if not os.path.exists(logo_path):
        logo_path = obter_caminho_recurso("logocura.png")
        
    if os.path.exists(logo_path):
        try:
            img = OpenpyxlImage(logo_path)
            img.width = 54
            img.height = 54
            ws.add_image(img, 'A1')
        except Exception:
            pass

    for col in range(1, 7):
        ws.cell(row=2, column=col).fill = PatternFill(start_color=COR_VERDE_FAIXA, end_color=COR_VERDE_FAIXA, fill_type="solid")

    cabecalhos = [
        "Ministro",
        "Código do Crachá",
        "Hora",
        "Data",
        "Quantidade de Ministros Hoje",
        "Total de Ministros"
    ]

    borda_fina = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    for idx, texto in enumerate(cabecalhos, 1):
        c = ws.cell(row=3, column=idx)
        c.value = texto
        c.font = Font(name="Segoe UI", size=11, bold=True, color="000000")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda_fina

    linha_atual = 4
    for reg in registros:
        nome, codigo, data_hora = reg
        try:
            dt_obj = datetime.strptime(data_hora, '%d/%m/%Y às %H:%M:%S')
            hora_str = dt_obj.strftime('%H:%M:%S')
            data_str = dt_obj.strftime('%d/%m/%Y')
        except Exception:
            partes = data_hora.split(" às ")
            data_str = partes[0] if len(partes) > 0 else data_hora
            hora_str = partes[1] if len(partes) > 1 else ""

        cor_fundo = COR_VERDE_ZEBRA if (linha_atual % 2 == 0) else COR_BRANCO
        fill_zebrada = PatternFill(start_color=cor_fundo, end_color=cor_fundo, fill_type="solid")

        ws.cell(row=linha_atual, column=1, value=nome)
        ws.cell(row=linha_atual, column=2, value=str(codigo))
        ws.cell(row=linha_atual, column=3, value=hora_str)
        ws.cell(row=linha_atual, column=4, value=data_str)
        
        if linha_atual == 4:
            ws.cell(row=linha_atual, column=5, value=total_hoje)
            ws.cell(row=linha_atual, column=6, value=total_cadastrados)

        for col in range(1, 7):
            cell = ws.cell(row=linha_atual, column=col)
            cell.fill = fill_zebrada
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = borda_fina
            cell.font = Font(name="Segoe UI", size=10)

        ws.row_dimensions[linha_atual].height = 22
        linha_atual += 1

    larguras = {1: 24, 2: 18, 3: 16, 4: 16, 5: 30, 6: 24}
    for col_idx, width in larguras.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    wb.save(caminho_final)
    return caminho_final

# ==========================================
# SUPORTE A CAMINHOS (.PY E .EXE)
# ==========================================
def obter_caminho_recurso(nome_arquivo):
    if getattr(sys, 'frozen', False):
        base_path = sys._MEIPASS
    else:
        base_path = os.path.dirname(os.path.abspath(__file__))
    
    caminho = os.path.join(base_path, nome_arquivo)
    if os.path.exists(caminho):
        return caminho
    if os.path.exists(nome_arquivo):
        return nome_arquivo
    return nome_arquivo

def obter_pasta_desktop():
    usuario = os.environ.get("USERPROFILE", os.path.expanduser("~"))
    desktop_onedrive = os.path.join(usuario, "OneDrive", "Desktop")
    desktop_onedrive_pt = os.path.join(usuario, "OneDrive", "Área de Trabalho")
    desktop_padrao = os.path.join(usuario, "Desktop")
    desktop_padrao_pt = os.path.join(usuario, "Área de Trabalho")

    for d in [desktop_onedrive, desktop_onedrive_pt, desktop_padrao, desktop_padrao_pt]:
        if os.path.exists(d):
            return d
    return desktop_padrao

# ==========================================
# BANCO DE DADOS
# ==========================================
def inicializar_banco():
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS funcionarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            codigo_barra TEXT UNIQUE NOT NULL
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS registros (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            funcionario_id INTEGER,
            data_hora DATETIME NOT NULL,
            FOREIGN KEY (funcionario_id) REFERENCES funcionarios(id)
        )
    ''')
    conn.commit()
    conn.close()

def registrar_leitura(codigo_lido):
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('SELECT id, nome FROM funcionarios WHERE codigo_barra = ?', (codigo_lido,))
    ministro = cursor.fetchone()
    
    if ministro:
        agora = datetime.now().strftime('%d/%m/%Y às %H:%M:%S')
        cursor.execute('INSERT INTO registros (funcionario_id, data_hora) VALUES (?, ?)', (ministro[0], agora))
        conn.commit()
        conn.close()
        return True, f"✓  {ministro[1]} — Ponto registrado ({agora})"
    else:
        conn.close()
        return False, f"✕  Código {codigo_lido} não encontrado no sistema"

def cadastrar_ministro(nome, codigo_barra):
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('INSERT INTO funcionarios (nome, codigo_barra) VALUES (?, ?)', (nome, codigo_barra))
    conn.commit()
    conn.close()

def listar_ministros():
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('SELECT id, nome, codigo_barra FROM funcionarios ORDER BY nome ASC')
    dados = cursor.fetchall()
    conn.close()
    return dados

def atualizar_nome_ministro(codigo_barra, novo_nome):
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('UPDATE funcionarios SET nome = ? WHERE codigo_barra = ?', (novo_nome, codigo_barra))
    conn.commit()
    conn.close()

def remover_ministro_banco(codigo_barra):
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('SELECT id, nome FROM funcionarios WHERE codigo_barra = ?', (codigo_barra,))
    ministro = cursor.fetchone()
    if ministro:
        func_id, nome = ministro
        cursor.execute('DELETE FROM registros WHERE funcionario_id = ?', (func_id,))
        cursor.execute('DELETE FROM funcionarios WHERE id = ?', (func_id,))
        conn.commit()
        conn.close()
        return True, nome
    conn.close()
    return False, None

def obter_estatisticas():
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('SELECT COUNT(*) FROM funcionarios')
    total_ministros = cursor.fetchone()[0]
    
    hoje = datetime.now().strftime('%d/%m/%Y')
    cursor.execute("SELECT COUNT(DISTINCT funcionario_id) FROM registros WHERE data_hora LIKE ?", (f"{hoje}%",))
    total_pontos_hoje = cursor.fetchone()[0]
    conn.close()
    return total_ministros, total_pontos_hoje

def obter_ultimos_registros(limite=15):
    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT f.nome, f.codigo_barra, r.data_hora
        FROM registros r
        JOIN funcionarios f ON r.funcionario_id = f.id
        ORDER BY r.id DESC
        LIMIT ?
    ''', (limite,))
    dados = cursor.fetchall()
    conn.close()
    return dados

# ==========================================
# EXPORTAÇÃO PERSONALIZADA (edição#n_dd-mm-aaaa)
# ==========================================
def exportar_excel_personalizado(texto_edicao=""):
    desktop = obter_pasta_desktop()
    pasta_destino = os.path.join(desktop, "relatorios")
    if not os.path.exists(pasta_destino):
        os.makedirs(pasta_destino)

    # Limpeza para evitar caracteres proibidos no Windows
    tag_edicao = texto_edicao.strip() if texto_edicao.strip() else "edição#"
    tag_edicao_limpa = "".join(c for c in tag_edicao if c.isalnum() or c in ('#', '_', '-'))
    
    data_arquivo = datetime.now().strftime('%d-%m-%Y')
    nome_arquivo = f"{tag_edicao_limpa}_{data_arquivo}.xlsx"
    caminho_final = os.path.join(pasta_destino, nome_arquivo)

    conn = sqlite3.connect('controle_ponto.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT f.nome, f.codigo_barra, r.data_hora
        FROM registros r
        JOIN funcionarios f ON r.funcionario_id = f.id
        ORDER BY r.id ASC
    ''')
    registros = cursor.fetchall()
    conn.close()

    total_cadastrados, total_hoje = obter_estatisticas()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controle de Ponto"
    ws.views.sheetView[0].showGridLines = True

    COR_VERDE_BANNER = "1F4E24"
    COR_VERDE_FAIXA  = "2E933C"
    COR_VERDE_ZEBRA  = "B8DABA"
    COR_BRANCO       = "FFFFFF"

    ws.row_dimensions[1].height = 48
    ws.row_dimensions[2].height = 6
    ws.row_dimensions[3].height = 28

    # Título do Banner
    ws.merge_cells('C1:F1')
    titulo_cell = ws['C1']
    titulo_cell.value = f"Controle de ponto do centro de cura ({tag_edicao})"
    titulo_cell.font = Font(name="Segoe UI", size=15, bold=True, color=COR_BRANCO)
    titulo_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 7):
        ws.cell(row=1, column=col).fill = PatternFill(start_color=COR_VERDE_BANNER, end_color=COR_VERDE_BANNER, fill_type="solid")

    logo_path = obter_caminho_recurso("logocura2.png")
    if not os.path.exists(logo_path):
        logo_path = obter_caminho_recurso("logocura.png")
        
    if os.path.exists(logo_path):
        try:
            img = OpenpyxlImage(logo_path)
            img.width = 54
            img.height = 54
            ws.add_image(img, 'A1')
        except Exception:
            pass

    for col in range(1, 7):
        ws.cell(row=2, column=col).fill = PatternFill(start_color=COR_VERDE_FAIXA, end_color=COR_VERDE_FAIXA, fill_type="solid")

    cabecalhos = [
        "Ministro",
        "Codigo do cracha",
        "Hora",
        "Data",
        "Quantidade de minsitro de hoje",
        "Quantidade de minsitros"
    ]

    borda_fina = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    for idx, texto in enumerate(cabecalhos, 1):
        c = ws.cell(row=3, column=idx)
        c.value = texto
        c.font = Font(name="Segoe UI", size=11, bold=False, color="000000")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda_fina

    linha_atual = 4
    for reg in registros:
        nome, codigo, data_hora = reg
        try:
            dt_obj = datetime.strptime(data_hora, '%d/%m/%Y às %H:%M:%S')
            hora_str = dt_obj.strftime('%H:%M:%S')
            data_str = dt_obj.strftime('%d/%m/%Y')
        except Exception:
            partes = data_hora.split(" às ")
            data_str = partes[0] if len(partes) > 0 else data_hora
            hora_str = partes[1] if len(partes) > 1 else ""

        cor_fundo = COR_VERDE_ZEBRA if (linha_atual % 2 == 0) else COR_BRANCO
        fill_zebrada = PatternFill(start_color=cor_fundo, end_color=cor_fundo, fill_type="solid")

        ws.cell(row=linha_atual, column=1, value=nome)
        ws.cell(row=linha_atual, column=2, value=str(codigo))
        ws.cell(row=linha_atual, column=3, value=hora_str)
        ws.cell(row=linha_atual, column=4, value=data_str)
        
        if linha_atual == 4:
            ws.cell(row=linha_atual, column=5, value=total_hoje)
            ws.cell(row=linha_atual, column=6, value=total_cadastrados)

        for col in range(1, 7):
            cell = ws.cell(row=linha_atual, column=col)
            cell.fill = fill_zebrada
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = borda_fina
            cell.font = Font(name="Segoe UI", size=10)

        ws.row_dimensions[linha_atual].height = 22
        linha_atual += 1

    larguras = {1: 22, 2: 18, 3: 16, 4: 16, 5: 28, 6: 24}
    for col_idx, width in larguras.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    wb.save(caminho_final)
    return caminho_final

def gerar_imagem_barcode(codigo_6_digitos, nome):
    folder = "crachas"
    if not os.path.exists(folder):
        os.makedirs(folder)
    
    barcode_class = barcode.get_barcode_class('code128')
    bar = barcode_class(codigo_6_digitos, writer=ImageWriter())
    
    safe_nome = "".join(c for c in nome if c.isalnum() or c in (' ', '_', '-')).rstrip()
    full_path = os.path.join(folder, f"{safe_nome}_{codigo_6_digitos}")
    saved_path = bar.save(full_path)
    return saved_path, codigo_6_digitos


# ==========================================
# INTERFACE GRÁFICA
# ==========================================
if __name__ == "__main__":
    inicializar_banco()
    
    ctk.set_appearance_mode("light")

    app = ctk.CTk()
    app.title("Centro de Cura — Sistema de Controle de Ponto")
    app.state("zoomed")
    
    # 1. Troca o ícone padrão pelo ícone da sua aplicação
    caminho_ico = obter_caminho_recurso("logocura.ico")
    if os.path.exists(caminho_ico):
        app.iconbitmap(caminho_ico)

    # 2. Pinta a barra nativa do Windows
    aplicar_tema_barra_titulo(app, cor_hex="#F5F7F4")
    
    COR_VERDE_PRI = "#0B4D3B"
    COR_VERDE_SEC = "#237A57"
    COR_FUNDO     = "#F5F7F4"
    COR_BRANCO    = "#FFFFFF"
    COR_BORDA     = "#DDE3DF"
    COR_BORDA_INP = "#C5D0CA"
    COR_TEXTO     = "#17221D"
    COR_CINZA     = "#66736C"
    COR_ERRO      = "#C94A4A"
    COR_ERRO_BG   = "#FDF0F0"
    COR_SUCESSO_BG= "#EDF7F3"

    app.configure(fg_color=COR_FUNDO)

    # 1. CABEÇALHO
    header = ctk.CTkFrame(app, fg_color=COR_BRANCO, height=72, corner_radius=0)
    header.pack(fill="x", side="top")
    header.pack_propagate(False)

    ctk.CTkFrame(app, fg_color=COR_BORDA, height=1, corner_radius=0).pack(fill="x", side="top")

    header_left = ctk.CTkFrame(header, fg_color="transparent")
    header_left.pack(side="left", padx=28)

    caminho_logo = obter_caminho_recurso("logocura2.png")
    if not os.path.exists(caminho_logo):
        caminho_logo = obter_caminho_recurso("logocura.png")

    if os.path.exists(caminho_logo):
        try:
            img_pil = Image.open(caminho_logo)
            logo_ctk = ctk.CTkImage(light_image=img_pil, dark_image=img_pil, size=(42, 42))
            lbl_logo = ctk.CTkLabel(header_left, image=logo_ctk, text="")
            lbl_logo.pack(side="left", padx=(0, 14))
        except Exception:
            pass

    title_box = ctk.CTkFrame(header_left, fg_color="transparent")
    title_box.pack(side="left")

    ctk.CTkLabel(title_box, text="CENTRO DE CURA", font=("Segoe UI", 18, "bold"), text_color=COR_TEXTO, anchor="w").pack(fill="x")
    ctk.CTkLabel(title_box, text="Sistema de Controle de Ponto", font=("Segoe UI", 12), text_color=COR_CINZA, anchor="w").pack(fill="x")

    def acao_exportar():
        edicao_texto = entry_edicao.get().strip()
        caminho_salvo = exportar_excel_personalizado(edicao_texto)
        messagebox.showinfo("Exportação Concluída", f"Planilha salva com sucesso em:\n\n{caminho_salvo}")

    btn_exportar_topo = ctk.CTkButton(
        header, text="📊  Exportar Excel", command=acao_exportar,
        height=38, width=150, corner_radius=10,
        fg_color=COR_BRANCO, hover_color=COR_FUNDO,
        border_width=1, border_color=COR_BORDA_INP,
        text_color=COR_VERDE_PRI, font=("Segoe UI", 13, "bold")
    )
    btn_exportar_topo.pack(side="right", padx=28)

    title_box = ctk.CTkFrame(header_left, fg_color="transparent")
    title_box.pack(side="left")

    ctk.CTkLabel(title_box, text="CENTRO DE CURA", font=("Segoe UI", 18, "bold"), text_color=COR_TEXTO, anchor="w").pack(fill="x")
    ctk.CTkLabel(title_box, text="Sistema de Controle de Ponto", font=("Segoe UI", 12), text_color=COR_CINZA, anchor="w").pack(fill="x")

    def acao_exportar():
        edicao_texto = entry_edicao.get().strip()
        caminho_salvo = exportar_excel_personalizado(edicao_texto)
        messagebox.showinfo("Exportação Concluída", f"Planilha salva com sucesso em:\n\n{caminho_salvo}")

    btn_exportar_topo = ctk.CTkButton(
        header, text="📊  Exportar Excel", command=acao_exportar,
        height=38, width=150, corner_radius=10,
        fg_color=COR_BRANCO, hover_color=COR_FUNDO,
        border_width=1, border_color=COR_BORDA_INP,
        text_color=COR_VERDE_PRI, font=("Segoe UI", 13, "bold")
    )
    btn_exportar_topo.pack(side="right", padx=28)

    # 2. WORKSPACE
    workspace = ctk.CTkFrame(app, fg_color=COR_FUNDO, corner_radius=0)
    workspace.pack(fill="both", expand=True)

    # 3. ABAS SUPERIORES
    nav_bar = ctk.CTkFrame(workspace, fg_color=COR_BRANCO, corner_radius=12, border_width=1, border_color=COR_BORDA)
    nav_bar.pack(pady=(28, 20))

    tab_btns = {}
    abas_topo = [
        ("inicio", "🏠 Início"),
        ("ministros", "👤 Ministros"),
        ("ponto", "🕒 Ponto"),
        ("relatorios", "📊 Relatórios"),
        ("config", "⚙ Config.")
    ]
    for col, (chave, texto) in enumerate(abas_topo):
        btn = ctk.CTkButton(
            nav_bar, text=texto, width=130, height=36, corner_radius=10,
            font=("Segoe UI", 12, "bold"), command=lambda c=chave: navegar_para(c)
        )
        btn.grid(row=0, column=col, padx=4, pady=4)
        tab_btns[chave] = btn

    # 4. CARD CENTRAL
    card_central = ctk.CTkFrame(
        workspace, width=540, height=520,
        fg_color=COR_BRANCO, corner_radius=18,
        border_width=1, border_color=COR_BORDA
    )
    card_central.pack(pady=(0, 20))
    card_central.pack_propagate(False)

    # ==========================================
    # TELAS
    # ==========================================

    # --- TELA 1: INÍCIO ---
    view_inicio = ctk.CTkFrame(card_central, fg_color="transparent")
    ctk.CTkLabel(view_inicio, text="PAINEL PRINCIPAL", font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO).pack(pady=(35, 6))
    ctk.CTkLabel(view_inicio, text="Visão geral do controle de presença", font=("Segoe UI", 13), text_color=COR_CINZA).pack(pady=(0, 25))

    cards_stat = ctk.CTkFrame(view_inicio, fg_color="transparent")
    cards_stat.pack(pady=10)

    card1 = ctk.CTkFrame(cards_stat, width=180, height=110, corner_radius=12, fg_color=COR_FUNDO, border_width=1, border_color=COR_BORDA)
    card1.grid(row=0, column=0, padx=12)
    card1.pack_propagate(False)
    lbl_stat_ministros = ctk.CTkLabel(card1, text="0", font=("Segoe UI", 26, "bold"), text_color=COR_VERDE_PRI)
    lbl_stat_ministros.pack(pady=(18, 2))
    ctk.CTkLabel(card1, text="Ministros Cadastrados", font=("Segoe UI", 11), text_color=COR_CINZA).pack()

    card2 = ctk.CTkFrame(cards_stat, width=180, height=110, corner_radius=12, fg_color=COR_FUNDO, border_width=1, border_color=COR_BORDA)
    card2.grid(row=0, column=1, padx=12)
    card2.pack_propagate(False)
    lbl_stat_pontos = ctk.CTkLabel(card2, text="0", font=("Segoe UI", 26, "bold"), text_color=COR_VERDE_PRI)
    lbl_stat_pontos.pack(pady=(18, 2))
    ctk.CTkLabel(card2, text="Presenças Hoje", font=("Segoe UI", 11), text_color=COR_CINZA).pack()

    btn_ir_leitura = ctk.CTkButton(
        view_inicio, text="Iniciar Leitura de Ponto ➔", command=lambda: navegar_para("ponto"),
        width=380, height=44, corner_radius=10, fg_color=COR_VERDE_PRI, hover_color=COR_VERDE_SEC,
        text_color=COR_BRANCO, font=("Segoe UI", 14, "bold")
    )
    btn_ir_leitura.pack(pady=(45, 0))

    # --- TELA 2: MINISTROS ---
    view_ministros = ctk.CTkFrame(card_central, fg_color="transparent")
    ctk.CTkLabel(view_ministros, text="CADASTRAR MINISTRO", font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO).pack(pady=(30, 4))
    ctk.CTkLabel(view_ministros, text="Gere crachás no padrão Code 128 com 6 dígitos", font=("Segoe UI", 13), text_color=COR_CINZA).pack(pady=(0, 20))

    form_box = ctk.CTkFrame(view_ministros, fg_color="transparent", width=360)
    form_box.pack()

    ctk.CTkLabel(form_box, text="Nome do ministro", font=("Segoe UI", 13, "bold"), text_color=COR_TEXTO, anchor="w").pack(fill="x", pady=(0, 5))
    entry_nome = ctk.CTkEntry(form_box, width=360, height=40, corner_radius=10, fg_color=COR_BRANCO, border_width=1, border_color=COR_BORDA_INP, font=("Segoe UI", 14), placeholder_text="Digite o nome completo", text_color=COR_TEXTO)
    entry_nome.pack(pady=(0, 14))

    ctk.CTkLabel(form_box, text="Código do crachá (Exatamente 6 dígitos numéricos)", font=("Segoe UI", 13, "bold"), text_color=COR_TEXTO, anchor="w").pack(fill="x", pady=(0, 5))
    entry_base = ctk.CTkEntry(form_box, width=360, height=40, corner_radius=10, fg_color=COR_BRANCO, border_width=1, border_color=COR_BORDA_INP, font=("Segoe UI", 14), placeholder_text="Ex: 123456", text_color=COR_TEXTO)
    entry_base.pack(pady=(0, 18))

    alerta_cad = ctk.CTkFrame(form_box, width=360, height=36, corner_radius=8, fg_color="transparent")
    alerta_cad.pack_propagate(False)
    lbl_alerta_cad = ctk.CTkLabel(alerta_cad, text="", font=("Segoe UI", 12, "bold"))
    lbl_alerta_cad.place(relx=0.5, rely=0.5, anchor="center")

    def salvar_cadastro():
        nome = entry_nome.get().strip()
        base = entry_base.get().strip()
        if not nome or not base.isdigit() or len(base) != 6:
            alerta_cad.configure(fg_color=COR_ERRO_BG, border_width=1, border_color=COR_ERRO)
            lbl_alerta_cad.configure(text="✕ Digite o nome e um código com exatamente 6 números", text_color=COR_ERRO)
            alerta_cad.pack(pady=(0, 10))
            return
        try:
            img_path, cod = gerar_imagem_barcode(base, nome)
            cadastrar_ministro(nome, cod)
            alerta_cad.configure(fg_color=COR_SUCESSO_BG, border_width=1, border_color=COR_VERDE_SEC)
            lbl_alerta_cad.configure(text=f"✓ Crachá {cod} gerado na pasta crachas/", text_color=COR_VERDE_PRI)
            alerta_cad.pack(pady=(0, 10))
            entry_nome.delete(0, 'end')
            entry_base.delete(0, 'end')
        except sqlite3.IntegrityError:
            alerta_cad.configure(fg_color=COR_ERRO_BG, border_width=1, border_color=COR_ERRO)
            lbl_alerta_cad.configure(text="✕ Este código de 6 dígitos já está em uso", text_color=COR_ERRO)
            alerta_cad.pack(pady=(0, 10))
        except Exception as e:
            messagebox.showerror("Erro", f"Ocorreu um erro: {e}")

    btn_cadastrar = ctk.CTkButton(form_box, text="Gerar crachá", command=salvar_cadastro, width=360, height=42, corner_radius=10, fg_color=COR_VERDE_PRI, hover_color=COR_VERDE_SEC, text_color=COR_BRANCO, font=("Segoe UI", 14, "bold"))
    btn_cadastrar.pack(pady=(5, 0))

    # --- TELA 3: PONTO COM EDIÇÃO# ---
    view_ponto = ctk.CTkFrame(card_central, fg_color="transparent")
    ctk.CTkLabel(view_ponto, text="REGISTRAR PONTO", font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO).pack(pady=(22, 4))
    ctk.CTkLabel(view_ponto, text="Passe o seu crachá", font=("Segoe UI", 13), text_color=COR_CINZA).pack(pady=(0, 12))

    icone_box = ctk.CTkFrame(view_ponto, width=88, height=54, corner_radius=12, fg_color=COR_FUNDO, border_width=1, border_color=COR_BORDA)
    icone_box.pack(pady=(0, 8))
    icone_box.pack_propagate(False)
    ctk.CTkLabel(icone_box, text="◉", font=("Segoe UI", 22), text_color=COR_VERDE_PRI).place(relx=0.5, rely=0.5, anchor="center")

    ctk.CTkLabel(view_ponto, text="Aproxime o crachá do leitor USB", font=("Segoe UI", 12), text_color=COR_CINZA).pack(pady=(0, 8))

    entry_leitura = ctk.CTkEntry(view_ponto, width=340, height=42, corner_radius=10, fg_color=COR_FUNDO, border_width=1, border_color=COR_VERDE_PRI, justify="center", font=("Segoe UI", 15), placeholder_text="Aguardando leitura...", text_color=COR_TEXTO)
    entry_leitura.pack(pady=(0, 8))

    # CAMPO DE EDIÇÃO (Prefixado com edição#)
    box_edicao = ctk.CTkFrame(view_ponto, fg_color="transparent")
    box_edicao.pack(pady=(0, 8))
    
    ctk.CTkLabel(box_edicao, text="Identificação da Edição:", font=("Segoe UI", 11, "bold"), text_color=COR_CINZA).pack(side="left", padx=(0, 8))
    entry_edicao = ctk.CTkEntry(box_edicao, width=170, height=32, corner_radius=8, font=("Segoe UI", 12, "bold"), text_color=COR_VERDE_PRI, fg_color=COR_FUNDO, border_color=COR_BORDA_INP)
    entry_edicao.pack(side="left")
    entry_edicao.insert(0, "edição#")

    # Garante que o prefixo "edição#" não seja apagado acidentalmente
    def validar_prefixo_edicao(event=None):
        val = entry_edicao.get()
        if not val.startswith("edição#"):
            nums = "".join(c for c in val if c.isdigit())
            entry_edicao.delete(0, 'end')
            entry_edicao.insert(0, f"edição#{nums}")

    entry_edicao.bind('<KeyRelease>', validar_prefixo_edicao)

    alerta_ponto = ctk.CTkFrame(view_ponto, width=340, height=38, corner_radius=10, fg_color="transparent")
    alerta_ponto.pack(pady=(0, 6))
    alerta_ponto.pack_propagate(False)
    lbl_alerta_ponto = ctk.CTkLabel(alerta_ponto, text="", font=("Segoe UI", 12, "bold"))
    lbl_alerta_ponto.place(relx=0.5, rely=0.5, anchor="center")

    def on_enter_leitura(event):
        cod = entry_leitura.get().strip()
        if cod:
            entry_leitura.delete(0, 'end')
            sucesso, msg = registrar_leitura(cod)
            if sucesso:
                alerta_ponto.configure(fg_color=COR_SUCESSO_BG, border_width=1, border_color=COR_VERDE_SEC)
                lbl_alerta_ponto.configure(text=msg, text_color=COR_VERDE_PRI)
            else:
                alerta_ponto.configure(fg_color=COR_ERRO_BG, border_width=1, border_color=COR_ERRO)
                lbl_alerta_ponto.configure(text=msg, text_color=COR_ERRO)
        return "break"

    entry_leitura.bind('<Return>', on_enter_leitura)
    ctk.CTkLabel(view_ponto, text="●  Leitor pronto para leitura", font=("Segoe UI", 12), text_color=COR_VERDE_SEC).pack(side="bottom", pady=16)

    # --- TELA 4: RELATÓRIOS ---
    view_relatorios = ctk.CTkFrame(card_central, fg_color="transparent")
    ctk.CTkLabel(view_relatorios, text="HISTÓRICO DE LEITURAS", font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO).pack(pady=(28, 4))
    ctk.CTkLabel(view_relatorios, text="Últimos registros gravados localmente", font=("Segoe UI", 13), text_color=COR_CINZA).pack(pady=(0, 14))

    lista_box = ctk.CTkScrollableFrame(view_relatorios, width=420, height=260, corner_radius=10, fg_color=COR_FUNDO, border_width=1, border_color=COR_BORDA)
    lista_box.pack(pady=(0, 16))

    def atualizar_relatorio_visual():
        for w in lista_box.winfo_children():
            w.destroy()
        dados = obter_ultimos_registros(15)
        if not dados:
            ctk.CTkLabel(lista_box, text="Nenhum registro encontrado.", font=("Segoe UI", 12), text_color=COR_CINZA).pack(pady=35)
            return
        for nome, cod, data_hora in dados:
            item = ctk.CTkFrame(lista_box, fg_color=COR_BRANCO, corner_radius=8, height=38)
            item.pack(fill="x", padx=4, pady=3)
            ctk.CTkLabel(item, text=f"{nome} ({cod})", font=("Segoe UI", 12, "bold"), text_color=COR_TEXTO).pack(side="left", padx=10, pady=6)
            ctk.CTkLabel(item, text=data_hora, font=("Segoe UI", 11), text_color=COR_CINZA).pack(side="right", padx=10, pady=6)

    ctk.CTkButton(view_relatorios, text="📊 Baixar Planilha na Área de Trabalho", command=acao_exportar, width=420, height=40, corner_radius=10, fg_color=COR_VERDE_PRI, hover_color=COR_VERDE_SEC, text_color=COR_BRANCO, font=("Segoe UI", 13, "bold")).pack()

    # --- TELA 5: CONFIGURAÇÕES ---
    view_config = ctk.CTkFrame(card_central, fg_color="transparent")
    ctk.CTkLabel(view_config, text="GERENCIAR MINISTROS", font=("Segoe UI", 20, "bold"), text_color=COR_TEXTO).pack(pady=(22, 2))
    ctk.CTkLabel(view_config, text="Altere o nome ou remova um ministro cadastrado", font=("Segoe UI", 13), text_color=COR_CINZA).pack(pady=(0, 14))

    box_gestao = ctk.CTkFrame(view_config, width=440, corner_radius=12, fg_color=COR_FUNDO, border_width=1, border_color=COR_BORDA)
    box_gestao.pack(padx=20, pady=(0, 10), fill="x")

    ctk.CTkLabel(box_gestao, text="Selecione o ministro:", font=("Segoe UI", 12, "bold"), text_color=COR_TEXTO, anchor="w").pack(fill="x", padx=16, pady=(12, 4))
    
    combo_ministros = ctk.CTkComboBox(
        box_gestao, width=400, height=36, corner_radius=8,
        fg_color=COR_BRANCO, border_color=COR_BORDA_INP, text_color=COR_TEXTO,
        values=["Carregando..."]
    )
    combo_ministros.pack(padx=16, pady=(0, 12))

    ctk.CTkLabel(box_gestao, text="Novo nome (para alteração):", font=("Segoe UI", 12, "bold"), text_color=COR_TEXTO, anchor="w").pack(fill="x", padx=16, pady=(0, 4))
    entry_novo_nome = ctk.CTkEntry(
        box_gestao, width=400, height=36, corner_radius=8,
        fg_color=COR_BRANCO, border_color=COR_BORDA_INP, text_color=COR_TEXTO,
        placeholder_text="Digite o novo nome..."
    )
    entry_novo_nome.pack(padx=16, pady=(0, 14))

    box_botoes = ctk.CTkFrame(box_gestao, fg_color="transparent")
    box_botoes.pack(fill="x", padx=16, pady=(0, 16))

    mapa_ministros = {}

    def atualizar_combo_config():
        dados = listar_ministros()
        mapa_ministros.clear()
        valores = []
        for _, nome, cod in dados:
            rotulo = f"{nome} (Código: {cod})"
            valores.append(rotulo)
            mapa_ministros[rotulo] = (cod, nome)
        
        if valores:
            combo_ministros.configure(values=valores)
            combo_ministros.set(valores[0])
            entry_novo_nome.delete(0, 'end')
            entry_novo_nome.insert(0, mapa_ministros[valores[0]][1])
        else:
            combo_ministros.configure(values=["Nenhum ministro cadastrado"])
            combo_ministros.set("Nenhum ministro cadastrado")
            entry_novo_nome.delete(0, 'end')

    def ao_selecionar_combo(escolha):
        if escolha in mapa_ministros:
            entry_novo_nome.delete(0, 'end')
            entry_novo_nome.insert(0, mapa_ministros[escolha][1])

    combo_ministros.configure(command=ao_selecionar_combo)

    def executar_alteracao_nome():
        selecao = combo_ministros.get()
        if selecao not in mapa_ministros:
            messagebox.showwarning("Aviso", "Selecione um ministro válido.")
            return
        novo = entry_novo_nome.get().strip()
        if not novo:
            messagebox.showwarning("Aviso", "Informe o novo nome do ministro.")
            return
        cod, _ = mapa_ministros[selecao]
        atualizar_nome_ministro(cod, novo)
        messagebox.showinfo("Sucesso", f"Nome atualizado para: {novo}")
        atualizar_combo_config()

    def executar_remocao():
        selecao = combo_ministros.get()
        if selecao not in mapa_ministros:
            messagebox.showwarning("Aviso", "Selecione um ministro válido.")
            return
        cod, nome = mapa_ministros[selecao]
        confirma = messagebox.askyesno(
            "Confirmar Remoção",
            f"Tem a certeza de que deseja remover o ministro:\n\n{nome} (Código: {cod})?\n\nIsto apagará o registro e o histórico de ponto deste ministro."
        )
        if confirma:
            remover_ministro_banco(cod)
            messagebox.showinfo("Removido", f"Ministro {nome} foi removido com sucesso!")
            atualizar_combo_config()

    btn_mudar_nome = ctk.CTkButton(
        box_botoes, text="✏ Mudar Nome", command=executar_alteracao_nome,
        width=190, height=38, corner_radius=8,
        fg_color=COR_VERDE_PRI, hover_color=COR_VERDE_SEC,
        text_color=COR_BRANCO, font=("Segoe UI", 12, "bold")
    )
    btn_mudar_nome.pack(side="left", padx=(0, 8))

    btn_remover = ctk.CTkButton(
        box_botoes, text="🗑 Remover Ministro", command=executar_remocao,
        width=190, height=38, corner_radius=8,
        fg_color=COR_ERRO, hover_color="#a83232",
        text_color=COR_BRANCO, font=("Segoe UI", 12, "bold")
    )
    btn_remover.pack(side="right")

    # Controle de Navegação
    telas = {
        "inicio": view_inicio,
        "ministros": view_ministros,
        "ponto": view_ponto,
        "relatorios": view_relatorios,
        "config": view_config
    }

    def navegar_para(destino):
        for t in telas.values():
            t.pack_forget()

        for k, b in tab_btns.items():
            if k == destino:
                b.configure(fg_color=COR_VERDE_PRI, hover_color=COR_VERDE_SEC, text_color=COR_BRANCO)
            else:
                b.configure(fg_color="transparent", hover_color=COR_FUNDO, text_color=COR_CINZA)

        telas[destino].pack(fill="both", expand=True)

        if destino == "inicio":
            tot_m, tot_p = obter_estatisticas()
            lbl_stat_ministros.configure(text=str(tot_m))
            lbl_stat_pontos.configure(text=str(tot_p))
        elif destino == "ponto":
            entry_leitura.focus_set()
        elif destino == "ministros":
            entry_nome.focus_set()
        elif destino == "relatorios":
            atualizar_relatorio_visual()
        elif destino == "config":
            atualizar_combo_config()

    navegar_para("ponto")

    app.mainloop()